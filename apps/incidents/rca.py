"""Root Cause Analysis report — the Forensic Analyst's structured deliverable.

A FORENSIC_RCA response request (TicketSubtask) owns at most one RCAReport. The
report is split the way the analyst works:

  * captured here, as data — Section 1 (prefilled from the ticket), affected
    assets (4), timeline (5), root causes (6.2), indicators (7) and the
    remediation list (8);
  * written in Word — every narrative section. The data above is rendered into a
    DOCX draft that the analyst finishes and uploads back to the request.

So this data is the source of truth up to the draft and Word is always the last
step: regenerating a draft does not carry Word edits back.

Every write path here ends in :func:`record_edit`, which keeps the stale-draft
notice, the audit trail and the OPEN→IN_PROGRESS nudge from drifting apart.
"""

import csv
import hashlib
import io
from datetime import datetime
from io import BytesIO
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.db.models import Max
from django.utils import timezone
from docx import Document

from . import history
from .ioc_values import (
    CAT_DOMAIN, CAT_FILE_NAME, CAT_FILE_PATH, CAT_HASH, CAT_IP, CAT_URL,
    INVENTORY_CATEGORY_CHOICES,
)
from .models import (
    AnalystIOC, IOCReviewStatus, RCAAsset, RCAIndicator, RCAReport,
    RCATimelineEntry, Ticket, TicketSubtask,
)
from .rca_content import RCA_TEMPLATE_VERSION
from .reports import (
    _THAI_MONTHS_ABBR, _chk, _expand_docx_repeat_rows, _replace_placeholders,
    _report_ticket_id,
)
from .ti_platform import _is_duplicate, create_manual_iocs

RCA_TEMPLATE_PATH = (
    Path(__file__).resolve().parent / 'report_templates' / 'rca_report_template_v1.docx'
)

# Audit labels — one per section the analyst can change.
SECTION_GENERAL = 'ส่วนที่ 1 ข้อมูลทั่วไป'
SECTION_ASSETS = 'ส่วนที่ 4 ทรัพย์สินที่ถูกบุกรุก'
SECTION_TIMELINE = 'ส่วนที่ 5 ลำดับเวลาเหตุการณ์'
SECTION_ROOT_CAUSES = 'ส่วนที่ 6.2 Root Cause'
SECTION_INDICATORS = 'ส่วนที่ 7 IOC'
SECTION_RECOMMENDATIONS = 'ส่วนที่ 8 คำแนะนำการแก้ไข'

# The IOC Database takes only these categories; email and account indicators
# stay in the report (decision 2026-09).
PUSHABLE_CATEGORIES = frozenset(code for code, _ in INVENTORY_CATEGORY_CHOICES)

# TicketIOC category → RCA indicator category. A bare file name has no RCA table
# of its own, so it lands in 7.1 for the analyst to complete into a path.
_TICKET_IOC_CATEGORY = {
    CAT_HASH: RCAIndicator.CAT_HASH,
    CAT_IP: RCAIndicator.CAT_IP,
    CAT_DOMAIN: RCAIndicator.CAT_DOMAIN,
    CAT_URL: RCAIndicator.CAT_URL,
    CAT_FILE_PATH: RCAIndicator.CAT_FILE_PATH,
    CAT_FILE_NAME: RCAIndicator.CAT_FILE_PATH,
}

# Ticket severity → the RCA form's SIEM scale (Medium is "Moderate" there).
_SIEM_SEVERITY = {'Critical': 'Critical', 'High': 'High', 'Medium': 'Moderate', 'Low': 'Low'}


# ── Formatting ────────────────────────────────────────────────────────────── #

def thai_date(value):
    """``15 พ.ค. 2565`` — the RCA form's date style (full Buddhist year)."""
    if not value:
        return ''
    if timezone.is_aware(value):
        value = timezone.localtime(value)
    return f'{value.day} {_THAI_MONTHS_ABBR[value.month]} {value.year + 543}'


def thai_datetime(value):
    """``15 พ.ค. 2565 19:02 น.``"""
    if not value:
        return ''
    if timezone.is_aware(value):
        value = timezone.localtime(value)
    return f'{thai_date(value)} {value:%H:%M} น.'


def case_number(rca):
    """1.1 — the RCA's case number, sharing the parent ticket's number."""
    return _report_ticket_id(rca.subtask.ticket, kind='RCA')


def case_number_for(subtask):
    """The RCA case number from the request alone, before any report exists —
    for the workspace's Start screen."""
    return _report_ticket_id(subtask.ticket, kind='RCA')


def _examiner_label(user):
    """1.15 — ``name (department · phone)``, dropping whatever is missing."""
    if not user:
        return ''
    name = user.get_full_name() or user.username
    profile = getattr(user, 'profile', None)
    extras = [
        part for part in (
            getattr(profile, 'department', ''), getattr(profile, 'phone', ''),
        ) if part
    ]
    return f'{name} ({" · ".join(extras)})' if extras else name


def _host_ip(ticket):
    return ' / '.join(part for part in (ticket.device_name, ticket.ip_address) if part)


# ── Creation and Section 1 prefill ────────────────────────────────────────── #

def initial_section1(subtask):
    """Section 1 values derived from the parent ticket."""
    ticket = subtask.ticket
    # Same value as the Incident/Event report's importance row; the Ticket and
    # RCAReport importance choices share their stored values.
    importance = ticket.report_importance
    owner = ticket.asset_owner or ''
    if ticket.asset_owner_name:
        owner = f'{owner} ({ticket.asset_owner_name})' if owner else ticket.asset_owner_name
    return {
        'incident_name': ticket.incident_name or '',
        'first_occurrence': ticket.event_occurred_at,
        'detected_at': ticket.incident_datetime,
        'importance': importance,
        'siem_severity': _SIEM_SEVERITY.get(ticket.severity, ''),
        'ncsa_severity': ticket.ncsa_severity or '',
        'threat_category': ticket.detailed_issue or '',
        'assets_examined': _host_ip(ticket),
        'asset_type': ticket.asset_type or '',
        'affected_systems': _host_ip(ticket),
        'asset_owner': owner,
        'examiner': _examiner_label(subtask.assigned_to),
        'related_refs': _report_ticket_id(ticket),
    }


def get_or_create_rca(subtask, user=None):
    """The request's RCAReport, created and prefilled on first use.

    Returns ``(rca, created)``. Creation seeds one Section 4 row from the ticket
    and pulls the ticket's indicators, so the analyst starts from the case as
    SOC recorded it rather than a blank form.
    """
    if subtask.subtask_type != TicketSubtask.TYPE_FORENSIC_RCA:
        raise ValidationError('รายงาน RCA ใช้ได้กับคำขอประเภท Forensics / RCA เท่านั้น')
    existing = RCAReport.objects.filter(subtask=subtask).first()
    if existing is not None:
        return existing, False
    try:
        with transaction.atomic():
            rca = RCAReport.objects.create(
                subtask=subtask, updated_by=user, **initial_section1(subtask),
            )
            ticket = subtask.ticket
            if ticket.device_name or ticket.ip_address:
                RCAAsset.objects.create(
                    rca=rca, order=0,
                    host=ticket.device_name or ticket.ip_address,
                    ip=ticket.ip_address or '',
                    detail=ticket.operating_system or '',
                )
            _add_ticket_indicators(rca)
    except IntegrityError:
        # A concurrent first open won the race; theirs is the report.
        return RCAReport.objects.get(subtask=subtask), False
    return rca, True


def refresh_section1_from_ticket(rca, user):
    """Re-apply the ticket prefill to Section 1, discarding the analyst's edits
    to those fields (the ticket may have been corrected since).

    Reads the ticket fresh: the whole point is to pick up a correction made
    after this report object — and the ticket cached on it — was loaded.
    """
    subtask = TicketSubtask.objects.select_related(
        'ticket', 'assigned_to__profile',
    ).get(pk=rca.subtask_id)
    for name, value in initial_section1(subtask).items():
        setattr(rca, name, value)
    rca.save()
    record_edit(rca, user, SECTION_GENERAL, 'ดึงข้อมูลจาก Ticket อีกครั้ง')


# ── Edit bookkeeping ──────────────────────────────────────────────────────── #

def start_request(subtask, user):
    """Move an OPEN request to IN_PROGRESS — the analyst has started the report.

    Re-reads the row so a status changed elsewhere since this object was loaded
    is not overwritten. Returns True when it moved.
    """
    fresh = TicketSubtask.objects.get(pk=subtask.pk)
    if fresh.status != TicketSubtask.STATUS_OPEN:
        return False
    fresh.status = TicketSubtask.STATUS_IN_PROGRESS
    fresh.save(update_fields=['status', 'updated_at'])
    history.record_subtask_status_change(
        fresh, TicketSubtask.STATUS_OPEN, TicketSubtask.STATUS_IN_PROGRESS, user,
        source='rca',
    )
    subtask.status = fresh.status
    subtask.status_changed_at = fresh.status_changed_at
    return True


def start_rca(subtask, user):
    """Begin the RCA: create + prefill the report and move OPEN → IN_PROGRESS.

    The one action behind the workspace's "Start RCA" button. Idempotent — a
    request already In Progress just returns its report, so a double click or a
    manager pressing Start after the analyst does nothing twice.
    """
    with transaction.atomic():
        report, _created = get_or_create_rca(subtask, user)
        start_request(subtask, user)
    return report


def record_edit(rca, user, section_label, summary):
    """Stamp the report as edited, audit the change and start the request."""
    now = timezone.now()
    RCAReport.objects.filter(pk=rca.pk).update(updated_at=now, updated_by=user)
    rca.updated_at, rca.updated_by = now, user
    history.record_rca_change(rca.subtask, section_label, summary, user)
    start_request(rca.subtask, user)


# ── Indicators: pull from the ticket, push to the IOC Database ────────────── #

def _next_orders(rca):
    return {
        row['category']: (row['top'] or 0) + 1
        for row in rca.indicators.values('category').annotate(top=Max('order'))
    }


def _add_ticket_indicators(rca):
    """Add the ticket's indicators that the report does not hold yet."""
    ticket = rca.subtask.ticket
    present = set(rca.indicators.values_list('category', 'value'))
    orders = _next_orders(rca)
    new = []

    def add(category, value):
        value = (value or '').strip()
        if not value or (category, value) in present:
            return
        present.add((category, value))
        order = orders.get(category, 0)
        orders[category] = order + 1
        new.append(RCAIndicator(
            rca=rca, category=category, value=value[:500],
            source=RCAIndicator.SOURCE_TICKET, order=order,
        ))

    for ioc in ticket.iocs.all():
        category = _TICKET_IOC_CATEGORY.get(ioc.category)
        if category:
            add(category, ioc.value)
    # The ticket keeps related accounts in its own field, one per line.
    for line in (ticket.ioc_user or '').splitlines():
        add(RCAIndicator.CAT_ACCOUNT, line)
    if new:
        RCAIndicator.objects.bulk_create(new)
    return len(new)


def pull_ticket_iocs(rca, user):
    """Bring in indicators added to the ticket since the report was started."""
    added = _add_ticket_indicators(rca)
    if added:
        record_edit(rca, user, SECTION_INDICATORS, f'ดึง IOC จาก Ticket +{added} รายการ')
    return added


def classify_ioc_push(rca):
    """Bucket every indicator for the push preview.

    Buckets, checked in this order: ``excluded``, ``unsupported`` (email /
    account), ``from_ticket`` (already in the IOC Database through the ticket),
    ``already_pushed``, ``duplicate`` (exists elsewhere — never re-owned by this
    case), ``restore`` (a removed entry comes back) and ``create``.
    """
    buckets = {key: [] for key in (
        'create', 'restore', 'duplicate', 'already_pushed',
        'from_ticket', 'excluded', 'unsupported',
    )}
    for indicator in rca.indicators.select_related('pushed_ioc'):
        if indicator.excluded:
            bucket = 'excluded'
        elif indicator.category not in PUSHABLE_CATEGORIES:
            bucket = 'unsupported'
        elif indicator.source == RCAIndicator.SOURCE_TICKET:
            bucket = 'from_ticket'
        elif indicator.pushed_ioc_id and indicator.pushed_ioc.is_active:
            bucket = 'already_pushed'
        elif _is_duplicate(indicator.category, indicator.value):
            bucket = 'duplicate'
        elif AnalystIOC.objects.filter(
            category=indicator.category, ioc_detail=indicator.value, is_active=False,
        ).exists():
            bucket = 'restore'
        else:
            bucket = 'create'
        buckets[bucket].append(indicator)
    return buckets


def _push_entry(indicator):
    entry = {'category': indicator.category, 'value': indicator.value, 'file_name': ''}
    if indicator.category == RCAIndicator.CAT_HASH and indicator.label:
        # A hash's label starts with the file it was found as.
        entry['file_name'] = indicator.label.strip().splitlines()[0][:255]
    # Carry the analyst's context into the IOC Database note, but never
    # overwrite an annotation the indicator already has there.
    context = ' — '.join(part.strip() for part in (indicator.label, indicator.note) if part.strip())
    if context and not IOCReviewStatus.objects.filter(
        category=indicator.category, value=indicator.value,
    ).exists():
        entry['note'] = context
    return entry


def push_iocs(rca, user):
    """Send the report's new indicators to the IOC Database, linked to the case.

    Returns ``{outcome: count}`` over create/restore candidates plus the counts
    of everything skipped, for the confirmation message.
    """
    buckets = classify_ioc_push(rca)
    candidates = buckets['create'] + buckets['restore']
    outcomes = {'created': 0, 'restored': 0, 'duplicate': len(buckets['duplicate'])}
    with transaction.atomic():
        results = create_manual_iocs(
            [_push_entry(indicator) for indicator in candidates], user,
            source_subtask=rca.subtask,
        )
        for indicator, (_entry, outcome) in zip(candidates, results):
            outcomes[outcome] = outcomes.get(outcome, 0) + 1
            if outcome in ('created', 'restored'):
                record = AnalystIOC.objects.filter(
                    category=indicator.category, ioc_detail=indicator.value, is_active=True,
                ).first()
                RCAIndicator.objects.filter(pk=indicator.pk).update(pushed_ioc=record)
        if outcomes['created'] or outcomes['restored']:
            record_edit(
                rca, user, SECTION_INDICATORS,
                f'ส่งเข้า IOC Database: สร้างใหม่ {outcomes["created"]} / '
                f'กู้คืน {outcomes["restored"]} / ซ้ำ {outcomes["duplicate"]}',
            )
    for key in ('already_pushed', 'from_ticket', 'excluded', 'unsupported'):
        outcomes[key] = len(buckets[key])
    return outcomes


# ── Timeline CSV/TSV import ───────────────────────────────────────────────── #

TIMELINE_CSV_COLUMNS = (
    'datetime', 'datetime_end', 'host', 'event', 'evidence_file', 'evidence_line', 'excerpt',
)
TIMELINE_IMPORT_MAX_BYTES = 2 * 1024 * 1024
TIMELINE_IMPORT_MAX_ROWS = 30
_TIMELINE_LIMITS = {'host': 255, 'evidence_file': 255, 'evidence_line': 64}


def timeline_csv_template():
    """The import template: header plus one example row (UTF-8, no BOM — the
    view adds one so Excel shows the Thai text)."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(TIMELINE_CSV_COLUMNS)
    writer.writerow([
        '2026-09-09 13:09', '', 'web-01 192.0.2.10',
        'ตัวอย่าง: ผู้ดูแลระงับบัญชีเว็บไซต์', 'system.log', '130',
        'ข้อความตัวอย่างจากไฟล์หลักฐาน',
    ])
    return buffer.getvalue()


def _decode(raw):
    # Excel's "CSV UTF-8" carries a BOM; its plain "CSV" on Thai Windows is cp874.
    for encoding in ('utf-8-sig', 'cp874'):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValidationError('อ่านไฟล์ไม่ได้ — กรุณาบันทึกเป็น CSV UTF-8')


def _parse_timeline_datetime(text):
    """``YYYY-MM-DD HH:MM[:SS]`` (Gregorian); naive values are UTC+7.

    A Buddhist-era year (≥ 2400 — never a real evidence year) is converted,
    since Thai users type it by habit.
    """
    text = (text or '').strip()
    if not text:
        return None
    try:
        value = datetime.fromisoformat(text)
    except ValueError:
        raise ValidationError(f'รูปแบบวันเวลาไม่ถูกต้อง "{text}" (ใช้ YYYY-MM-DD HH:MM)')
    if value.year >= 2400:
        value = value.replace(year=value.year - 543)
    if timezone.is_naive(value):
        value = timezone.make_aware(value, timezone.get_default_timezone())
    return value


def import_timeline_csv(rca, uploaded_file, user):
    """Append the rows of a timeline CSV/TSV to Section 5.

    All-or-nothing: every row is validated first and any error rejects the whole
    file, so a half-imported timeline never has to be cleaned up by hand.
    Returns the number of rows added.
    """
    raw = uploaded_file.read(TIMELINE_IMPORT_MAX_BYTES + 1)
    if len(raw) > TIMELINE_IMPORT_MAX_BYTES:
        raise ValidationError('ไฟล์ใหญ่เกิน 2 MB')
    text = _decode(raw)
    first_line = text.split('\n', 1)[0]
    delimiter = '\t' if first_line.count('\t') > first_line.count(',') else ','
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    header = [cell.strip().lower() for cell in next(reader, [])]
    if tuple(header[:len(TIMELINE_CSV_COLUMNS)]) != TIMELINE_CSV_COLUMNS:
        raise ValidationError(
            'หัวคอลัมน์ไม่ตรงกับ template — ต้องเป็น: ' + ', '.join(TIMELINE_CSV_COLUMNS)
        )

    entries, errors = [], []
    for row in reader:
        if not any(cell.strip() for cell in row):
            continue
        if len(entries) >= TIMELINE_IMPORT_MAX_ROWS:
            errors.append(f'เกิน {TIMELINE_IMPORT_MAX_ROWS} แถวต่อไฟล์')
            break
        values = [cell.strip() for cell in row][:len(TIMELINE_CSV_COLUMNS)]
        values += [''] * (len(TIMELINE_CSV_COLUMNS) - len(values))
        cells = dict(zip(TIMELINE_CSV_COLUMNS, values))
        try:
            start = _parse_timeline_datetime(cells['datetime'])
            if start is None:
                raise ValidationError('ต้องระบุวันเวลา (datetime)')
            end = _parse_timeline_datetime(cells['datetime_end'])
            if end is not None and end < start:
                raise ValidationError('datetime_end ต้องไม่ก่อน datetime')
            if not cells['event']:
                raise ValidationError('ต้องระบุเหตุการณ์ (event)')
            for key, limit in _TIMELINE_LIMITS.items():
                if len(cells[key]) > limit:
                    raise ValidationError(f'{key} ยาวเกิน {limit} ตัวอักษร')
        except ValidationError as exc:
            errors.append(f'บรรทัด {reader.line_num}: {"; ".join(exc.messages)}')
            continue
        entries.append(RCATimelineEntry(
            rca=rca, occurred_at=start, occurred_until=end,
            host=cells['host'], event=cells['event'],
            evidence_file=cells['evidence_file'], evidence_line=cells['evidence_line'],
            excerpt=cells['excerpt'],
        ))

    if errors:
        shown = errors[:20]
        if len(errors) > 20:
            shown.append(f'…และอีก {len(errors) - 20} บรรทัด')
        raise ValidationError(shown)
    if not entries:
        raise ValidationError('ไม่พบแถวข้อมูลในไฟล์')
    with transaction.atomic():
        RCATimelineEntry.objects.bulk_create(entries)
        name = getattr(uploaded_file, 'name', '') or 'CSV'
        record_edit(rca, user, SECTION_TIMELINE, f'+{len(entries)} แถว (นำเข้าไฟล์ {name})')
    return len(entries)


# ── DOCX draft ────────────────────────────────────────────────────────────── #

def _timeline_when(entry):
    """The timeline's date column: ``21 เม.ย. 2566 16:27–16:38`` (a same-day span
    collapses to a time range; a cross-day one shows both dates)."""
    start = timezone.localtime(entry.occurred_at)
    text = f'{thai_date(entry.occurred_at)} {start:%H:%M}'
    if entry.occurred_until:
        end = timezone.localtime(entry.occurred_until)
        if end.date() == start.date():
            text += f'–{end:%H:%M}'
        else:
            text += f' – {thai_date(entry.occurred_until)} {end:%H:%M}'
    return text


def _evidence_cell(entry):
    parts = [entry.evidence_file.strip()] if entry.evidence_file.strip() else []
    if entry.evidence_line.strip():
        parts.append(f'line {entry.evidence_line.strip()}')
    return '\n'.join(parts)


def _first_occurrence_text(rca):
    """First-occurrence date, with the optional confirmable note appended."""
    base = thai_datetime(rca.first_occurrence)
    note = (rca.first_occurrence_note or '').strip()
    if base and note:
        return f'{base} ({note})'
    return base or note or '-'


def _scope_text(rca):
    """Scope period rendered as ``start – end`` (or whichever side is set)."""
    start = thai_datetime(rca.scope_start)
    end = thai_datetime(rca.scope_end)
    if start and end:
        return f'{start} – {end}'
    return start or end or '-'


def _section1_context(rca):
    """The flat ``{{rca_*}}`` values for Section 1 and the page header."""
    threat = dict(Ticket.DETAILED_ISSUE_CHOICES).get(rca.threat_category, rca.threat_category)
    forensic = set(rca.forensic_types or [])
    context = {
        'rca_case_no': case_number(rca),
        'rca_incident_name': rca.incident_name or '-',
        'rca_first_occurrence': _first_occurrence_text(rca),
        'rca_detected': thai_datetime(rca.detected_at) or '-',
        'rca_scope': _scope_text(rca),
        'rca_threat_category': threat or '-',
        'rca_assets_examined': rca.assets_examined or '-',
        'rca_affected_systems': rca.affected_systems or '-',
        'rca_asset_owner': rca.asset_owner or '-',
        'rca_examiner': rca.examiner or '-',
        'rca_related_refs': rca.related_refs or '-',
    }
    for key, _label in RCAReport.FORENSIC_TYPE_CHOICES:
        context[f'rca_chk_ft_{key}'] = _chk(key in forensic)
    context['rca_chk_imp_general'] = _chk(rca.importance == RCAReport.IMPORTANCE_GENERAL)
    context['rca_chk_imp_important'] = _chk(rca.importance == RCAReport.IMPORTANCE_IMPORTANT)
    context['rca_chk_imp_critical'] = _chk(rca.importance == RCAReport.IMPORTANCE_CRITICAL)
    for level in ('low', 'moderate', 'high', 'critical'):
        context[f'rca_chk_siem_{level}'] = _chk(rca.siem_severity == level.capitalize())
    context['rca_chk_ncsa_non_severe'] = _chk(rca.ncsa_severity == Ticket.NCSA_SEVERITY_NON_SEVERE)
    context['rca_chk_ncsa_severe'] = _chk(rca.ncsa_severity == Ticket.NCSA_SEVERITY_SEVERE)
    context['rca_chk_ncsa_critical'] = _chk(rca.ncsa_severity == Ticket.NCSA_SEVERITY_CRITICAL)
    context['rca_chk_asset_computer'] = _chk(rca.asset_type == 'Computer')
    context['rca_chk_asset_server'] = _chk(rca.asset_type == 'Server')
    context['rca_chk_asset_network'] = _chk(rca.asset_type == 'Network Device')
    context['rca_chk_asset_unknown'] = _chk(
        rca.asset_type not in ('Computer', 'Server', 'Network Device')
    )
    return context


def _repeat_rows(rca):
    """The item dicts for every repeatable table, keyed by placeholder prefix.

    Indicators fan out into their per-category tables; recommendations carry the
    derived ``(RC-n)`` codes of the root causes they address.
    """
    assets = [
        {'as_host': a.host, 'as_ip': a.ip, 'as_detail': a.detail}
        for a in rca.assets.all()
    ]
    timeline = [
        {
            'tl_when': _timeline_when(e), 'tl_host': e.host, 'tl_event': e.event,
            'tl_evidence': _evidence_cell(e), 'tl_excerpt': e.excerpt,
        }
        for e in rca.timeline.all()
    ]
    root_causes = list(rca.root_causes.all())
    code_by_pk = {rc.pk: f'RC-{index}' for index, rc in enumerate(root_causes, start=1)}
    rc_rows = [
        {
            'rc_code': code_by_pk[rc.pk], 'rc_category': rc.category, 'rc_cause': rc.cause,
            'rc_evidence': rc.evidence_ref, 'rc_excerpt': rc.excerpt,
        }
        for rc in root_causes
    ]
    indicators = list(rca.indicators.all())

    def by_cat(*categories):
        return [ind for ind in indicators if ind.category in categories]

    file_paths = [
        {'fp_value': i.value, 'fp_host': i.host, 'fp_note': i.note or i.label}
        for i in by_cat(RCAIndicator.CAT_FILE_PATH)
    ]
    ips = [
        {'ip_value': i.value, 'ip_label': i.label, 'ip_note': i.note}
        for i in by_cat(RCAIndicator.CAT_IP)
    ]
    web = [
        {
            'web_value': i.value,
            'web_label': i.label or i.get_category_display(),
            'web_note': i.note,
        }
        for i in by_cat(RCAIndicator.CAT_URL, RCAIndicator.CAT_DOMAIN, RCAIndicator.CAT_EMAIL)
    ]
    hashes = [
        {'hash_label': i.label or i.host, 'hash_value': i.value}
        for i in by_cat(RCAIndicator.CAT_HASH)
    ]
    accounts = [
        {'acct_value': i.value, 'acct_host': i.host, 'acct_note': i.note}
        for i in by_cat(RCAIndicator.CAT_ACCOUNT)
    ]
    recommendations = []
    for index, rec in enumerate(rca.recommendations.all(), start=1):
        codes = [code_by_pk[rc.pk] for rc in rec.root_causes.all() if rc.pk in code_by_pk]
        action = rec.action
        if codes:
            action = f'{action} ({", ".join(codes)})'
        recommendations.append({'rec_no': str(index), 'rec_action': action})
    return {
        'as': assets, 'tl': timeline, 'rc': rc_rows, 'fp': file_paths, 'ip': ips,
        'web': web, 'hash': hashes, 'acct': accounts, 'rec': recommendations,
    }


def generate_rca_draft(rca, user):
    """Render the prefilled RCA draft DOCX and record its provenance.

    Fills Section 1 and clones each evidence table's prototype row per stored
    item, leaving the narrative sections' guidance for Word. Returns
    ``(filename, content_bytes)``. Records draft_* on the RCAReport only — never
    the ticket's report_* provenance (that belongs to the Incident/Event report).
    """
    doc = Document(str(RCA_TEMPLATE_PATH))
    for prefix, rows in _repeat_rows(rca).items():
        _expand_docx_repeat_rows(doc, prefix, rows)
    _replace_placeholders(doc, _section1_context(rca))

    buffer = BytesIO()
    doc.save(buffer)
    content = buffer.getvalue()
    digest = hashlib.sha256(content).hexdigest()
    now = timezone.now()
    RCAReport.objects.filter(pk=rca.pk).update(
        draft_generated_at=now, draft_generated_by=user,
        draft_sha256=digest, draft_template_version=RCA_TEMPLATE_VERSION,
    )
    rca.draft_generated_at, rca.draft_generated_by = now, user
    rca.draft_sha256, rca.draft_template_version = digest, RCA_TEMPLATE_VERSION

    filename = f'report_{case_number(rca)}_{RCA_TEMPLATE_VERSION}.docx'
    return filename, content
