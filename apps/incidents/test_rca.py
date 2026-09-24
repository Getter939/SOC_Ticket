import importlib.util
import re
import tempfile
from datetime import datetime
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.contrib.auth.models import User
from django.core import mail
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone
from docx import Document

from apps.accounts.models import UserProfile
from apps.accounts.testing import MFATestCase as TestCase

from . import rca as rca_service
from .models import (
    AnalystIOC,
    IOCReviewStatus,
    RCAAsset,
    RCAIndicator,
    RCARecommendation,
    RCAReport,
    RCARootCause,
    RCATimelineEntry,
    Ticket,
    TicketFieldChange,
    TicketIOC,
    TicketSubtask,
)
from .rca_content import RCA_FORENSIC_TYPES, RCA_REPEAT_TABLES, RCA_SECTION1_ROWS
from .reports import REPORT_CONTENT_TYPE, _iter_paragraphs

HASH_A = 'a' * 64
HASH_B = 'b' * 64


def _user(username, role, *, tier='', department='Test', phone='000', **names):
    user = User.objects.create_user(username=username, password='testpass123', **names)
    UserProfile.objects.create(
        user=user, role=role, tier=tier, department=department, phone=phone,
    )
    return user


def _ticket(**overrides):
    values = {
        'device_name': 'web-01',
        'ip_address': '192.0.2.10',
        'issue_description': 'RCA service contract',
        'classification': Ticket.CLASSIFICATION_INCIDENT,
    }
    values.update(overrides)
    return Ticket.objects.create(**values)


def _rca_request(ticket, assignee, **overrides):
    values = {
        'ticket': ticket,
        'subtask_type': TicketSubtask.TYPE_FORENSIC_RCA,
        'title': 'Root cause of the defacement',
        'assigned_to': assignee,
    }
    values.update(overrides)
    return TicketSubtask.objects.create(**values)


def _bkk(*args):
    return timezone.make_aware(datetime(*args), timezone.get_default_timezone())


def _set_ticket_status(ticket, status):
    Ticket.objects.filter(pk=ticket.pk).update(status=status)
    ticket.refresh_from_db()


class RCAPrefillTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.forensic = _user(
            'rca-forensic', UserProfile.ROLE_FORENSIC,
            department='ปปกก.', phone='02-574-8209',
            first_name='วศิน', last_name='เชาว์บวร',
        )

    def test_section_one_is_prefilled_from_the_ticket(self):
        ticket = _ticket(
            incident_name='Website defaced',
            severity='Medium',
            ncsa_severity=Ticket.NCSA_SEVERITY_SEVERE,
            detailed_issue='Root Intrusion',
            asset_type='Server',
            operating_system='Ubuntu 18.04',
            asset_owner='ชบดจ.',
            asset_owner_name='Somchai',
            event_occurred_at=_bkk(2022, 5, 15, 19, 2),
            incident_datetime=_bkk(2026, 9, 9, 8, 30),
        )
        subtask = _rca_request(ticket, self.forensic)

        report, created = rca_service.get_or_create_rca(subtask, self.forensic)

        self.assertTrue(created)
        self.assertEqual(report.incident_name, 'Website defaced')
        self.assertEqual(report.first_occurrence, _bkk(2022, 5, 15, 19, 2))
        self.assertEqual(report.detected_at, _bkk(2026, 9, 9, 8, 30))
        self.assertEqual(report.siem_severity, 'Moderate')
        self.assertEqual(report.ncsa_severity, Ticket.NCSA_SEVERITY_SEVERE)
        self.assertEqual(report.threat_category, 'Root Intrusion')
        self.assertEqual(report.importance, RCAReport.IMPORTANCE_IMPORTANT)
        self.assertEqual(report.asset_type, 'Server')
        self.assertEqual(report.assets_examined, 'web-01 / 192.0.2.10')
        self.assertEqual(report.asset_owner, 'ชบดจ. (Somchai)')
        self.assertEqual(report.examiner, 'วศิน เชาว์บวร (ปปกก. · 02-574-8209)')
        period = ticket.ticket_id[len('SOC-'):]
        self.assertEqual(report.related_refs, f'SOC-INC-{period}')
        self.assertEqual(rca_service.case_number(report), f'SOC-RCA-{period}')
        asset = report.assets.get()
        self.assertEqual(
            (asset.host, asset.ip, asset.detail), ('web-01', '192.0.2.10', 'Ubuntu 18.04'),
        )

    def test_second_open_returns_the_same_report(self):
        subtask = _rca_request(_ticket(), self.forensic)
        first, _ = rca_service.get_or_create_rca(subtask, self.forensic)

        second, created = rca_service.get_or_create_rca(subtask, self.forensic)

        self.assertFalse(created)
        self.assertEqual(first.pk, second.pk)

    def test_importance_follows_emergency_and_classification(self):
        cases = (
            ({'is_emergency': True}, RCAReport.IMPORTANCE_CRITICAL),
            ({'classification': Ticket.CLASSIFICATION_EVENT}, RCAReport.IMPORTANCE_GENERAL),
            # The analyst's pick wins over the legacy classification rule.
            ({'importance': Ticket.IMPORTANCE_CRITICAL}, RCAReport.IMPORTANCE_CRITICAL),
            ({'classification': Ticket.CLASSIFICATION_EVENT,
              'importance': Ticket.IMPORTANCE_IMPORTANT}, RCAReport.IMPORTANCE_IMPORTANT),
        )
        for overrides, expected in cases:
            with self.subTest(overrides=overrides):
                subtask = _rca_request(_ticket(**overrides), self.forensic)
                report, _ = rca_service.get_or_create_rca(subtask, self.forensic)
                self.assertEqual(report.importance, expected)

    def test_unknown_severity_leaves_the_siem_row_blank(self):
        subtask = _rca_request(_ticket(severity='Unknown'), self.forensic)

        report, _ = rca_service.get_or_create_rca(subtask, self.forensic)

        self.assertEqual(report.siem_severity, '')

    def test_refuses_a_non_rca_request(self):
        red_team = _user('rca-redteam', UserProfile.ROLE_REDTEAM_MANAGER)
        subtask = _rca_request(
            _ticket(), red_team, subtask_type=TicketSubtask.TYPE_VA_PT,
        )

        with self.assertRaises(ValidationError):
            rca_service.get_or_create_rca(subtask, red_team)

    def test_refresh_reapplies_the_ticket_values(self):
        ticket = _ticket(incident_name='Original name')
        subtask = _rca_request(ticket, self.forensic)
        report, _ = rca_service.get_or_create_rca(subtask, self.forensic)
        report.incident_name = 'Analyst edit'
        report.save()
        Ticket.objects.filter(pk=ticket.pk).update(incident_name='Corrected on ticket')

        rca_service.refresh_section1_from_ticket(report, self.forensic)

        report.refresh_from_db()
        self.assertEqual(report.incident_name, 'Corrected on ticket')


class RCATimelineImportTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.forensic = _user('rca-csv-forensic', UserProfile.ROLE_FORENSIC)

    def setUp(self):
        self.subtask = _rca_request(_ticket(), self.forensic)
        self.report, _ = rca_service.get_or_create_rca(self.subtask, self.forensic)

    def _import(self, text, *, encoding='utf-8-sig', name='timeline.csv'):
        upload = SimpleUploadedFile(name, text.encode(encoding))
        return rca_service.import_timeline_csv(self.report, upload, self.forensic)

    def test_comma_file_with_thai_text_and_a_time_range(self):
        header = ','.join(rca_service.TIMELINE_CSV_COLUMNS)
        text = (
            f'{header}\n'
            '2023-04-21 16:27,2023-04-21 16:38,NTDA,Brute Force MariaDB 2690 ครั้ง,'
            'authfail.log,14,Access denied for user\n'
            '2022-05-15 19:02:15,,ridcm,วาง Backdoor ALFA Shell,outputs.txt,63–64,"a, b"\n'
        )

        added = self._import(text)

        self.assertEqual(added, 2)
        first, second = RCATimelineEntry.objects.filter(rca=self.report)
        self.assertEqual(first.event, 'วาง Backdoor ALFA Shell')
        self.assertEqual(first.excerpt, 'a, b')
        self.assertEqual(timezone.localtime(first.occurred_at).hour, 19)
        self.assertEqual(second.evidence_line, '14')
        self.assertEqual(timezone.localtime(second.occurred_until).minute, 38)

    def test_tab_separated_file(self):
        header = '\t'.join(rca_service.TIMELINE_CSV_COLUMNS)
        text = f'{header}\n2026-09-09 13:09\t\tNTDA\tระงับบัญชี claim\tsystem.log\t130\t\n'

        self.assertEqual(self._import(text, name='timeline.tsv'), 1)

    def test_cp874_encoded_file(self):
        header = ','.join(rca_service.TIMELINE_CSV_COLUMNS)
        text = f'{header}\n2026-09-09 13:09,,NTDA,ระงับบัญชี,,,\n'

        self._import(text, encoding='cp874')

        self.assertEqual(self.report.timeline.get().event, 'ระงับบัญชี')

    def test_buddhist_year_is_converted(self):
        header = ','.join(rca_service.TIMELINE_CSV_COLUMNS)

        self._import(f'{header}\n2569-09-09 13:09,,,event,,,\n')

        self.assertEqual(timezone.localtime(self.report.timeline.get().occurred_at).year, 2026)

    def test_any_bad_row_rejects_the_whole_file(self):
        header = ','.join(rca_service.TIMELINE_CSV_COLUMNS)
        text = (
            f'{header}\n'
            '2026-09-09 13:09,,NTDA,fine row,,,\n'
            '09/09/2026,,NTDA,bad date,,,\n'
            '2026-09-09 14:00,2026-09-09 13:00,NTDA,ends before it starts,,,\n'
            '2026-09-09 15:00,,NTDA,,,,\n'
        )

        with self.assertRaises(ValidationError) as caught:
            self._import(text)

        messages = ' | '.join(caught.exception.messages)
        self.assertIn('บรรทัด 3', messages)
        self.assertIn('บรรทัด 4', messages)
        self.assertIn('บรรทัด 5', messages)
        self.assertFalse(RCATimelineEntry.objects.filter(rca=self.report).exists())

    def test_wrong_header_is_rejected(self):
        with self.assertRaises(ValidationError):
            self._import('when,what\n2026-09-09 13:09,x\n')

    def test_template_round_trips(self):
        template = rca_service.timeline_csv_template()

        self.assertTrue(template.startswith(','.join(rca_service.TIMELINE_CSV_COLUMNS)))
        self.assertEqual(self._import(template), 1)

    def test_import_is_audited_and_starts_the_request(self):
        header = ','.join(rca_service.TIMELINE_CSV_COLUMNS)

        self._import(f'{header}\n2026-09-09 13:09,,,event,,,\n')

        self.subtask.refresh_from_db()
        self.assertEqual(self.subtask.status, TicketSubtask.STATUS_IN_PROGRESS)
        self.assertTrue(TicketFieldChange.objects.filter(
            subtask=self.subtask, field_name='rca', new_value__contains='+1',
        ).exists())
        self.assertTrue(TicketFieldChange.objects.filter(
            subtask=self.subtask, field_name='status', source='rca',
        ).exists())

    def test_import_accepts_thirty_rows_and_rejects_thirty_one_atomically(self):
        self.assertEqual(rca_service.TIMELINE_IMPORT_MAX_ROWS, 30)
        header = ','.join(rca_service.TIMELINE_CSV_COLUMNS)
        rows = [
            f'2026-09-09 13:{minute:02d},,,event {minute},,,'
            for minute in range(31)
        ]

        self.assertEqual(self._import('\n'.join([header, *rows[:30]])), 30)
        self.report.timeline.all().delete()

        with self.assertRaises(ValidationError) as caught:
            self._import('\n'.join([header, *rows]))

        self.assertIn('เกิน 30 แถวต่อไฟล์', caught.exception.messages)
        self.assertFalse(self.report.timeline.exists())


class RCAIndicatorTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.forensic = _user('rca-ioc-forensic', UserProfile.ROLE_FORENSIC)

    def setUp(self):
        self.ticket = _ticket(ioc_user='ginfo\nclaim\n')
        TicketIOC.objects.create(ticket=self.ticket, category='ip', value='203.0.113.5')
        TicketIOC.objects.create(ticket=self.ticket, category='hash', value=HASH_A)
        TicketIOC.objects.create(ticket=self.ticket, category='file_name', value='js.php')
        self.subtask = _rca_request(self.ticket, self.forensic)
        self.report, _ = rca_service.get_or_create_rca(self.subtask, self.forensic)

    def _add(self, category, value, **extra):
        return RCAIndicator.objects.create(
            rca=self.report, category=category, value=value,
            source=RCAIndicator.SOURCE_ANALYST, **extra,
        )

    def test_ticket_indicators_are_pulled_on_creation(self):
        pulled = set(self.report.indicators.values_list('category', 'value', 'source'))

        self.assertEqual(pulled, {
            ('ip', '203.0.113.5', 'ticket'),
            ('hash', HASH_A, 'ticket'),
            ('file_path', 'js.php', 'ticket'),
            ('account', 'ginfo', 'ticket'),
            ('account', 'claim', 'ticket'),
        })

    def test_pulling_again_adds_only_new_ticket_indicators(self):
        self.assertEqual(rca_service.pull_ticket_iocs(self.report, self.forensic), 0)
        TicketIOC.objects.create(ticket=self.ticket, category='domain', value='evil.example')

        self.assertEqual(rca_service.pull_ticket_iocs(self.report, self.forensic), 1)
        self.assertEqual(self.report.indicators.count(), 6)

    def test_push_buckets_and_links_new_indicators_to_the_case(self):
        created = self._add(
            'ip', '198.51.100.7', label='เครือข่ายภายนอก', note='ส่งคำขอ CVE-2020-25213',
        )
        self._add('domain', 'responder.example', excluded=True)
        self._add('email', 'attacker@example.com')
        AnalystIOC.objects.create(ext_id='MAN-9001', category='hash', ioc_detail=HASH_B)
        self._add('hash', HASH_B)
        AnalystIOC.objects.create(
            ext_id='MAN-9002', category='url', ioc_detail='https://bad.example/x',
            is_active=False,
        )
        restored = self._add('url', 'https://bad.example/x')

        buckets = rca_service.classify_ioc_push(self.report)

        def values(key):
            return {indicator.value for indicator in buckets[key]}

        self.assertEqual(values('create'), {'198.51.100.7'})
        self.assertEqual(values('restore'), {'https://bad.example/x'})
        self.assertEqual(values('duplicate'), {HASH_B})
        self.assertEqual(values('excluded'), {'responder.example'})
        self.assertEqual(values('unsupported'), {'attacker@example.com', 'ginfo', 'claim'})
        self.assertEqual(values('from_ticket'), {'203.0.113.5', HASH_A, 'js.php'})

        outcomes = rca_service.push_iocs(self.report, self.forensic)

        self.assertEqual(outcomes['created'], 1)
        self.assertEqual(outcomes['restored'], 1)
        self.assertEqual(outcomes['duplicate'], 1)
        new_record = AnalystIOC.objects.get(category='ip', ioc_detail='198.51.100.7')
        self.assertTrue(new_record.ext_id.startswith('MAN-'))
        self.assertEqual(new_record.source_subtask, self.subtask)
        back = AnalystIOC.objects.get(ext_id='MAN-9002')
        self.assertTrue(back.is_active)
        self.assertEqual(back.source_subtask, self.subtask)
        self.assertIsNone(AnalystIOC.objects.get(ext_id='MAN-9001').source_subtask)
        created.refresh_from_db()
        restored.refresh_from_db()
        self.assertEqual(created.pushed_ioc, new_record)
        self.assertEqual(restored.pushed_ioc, back)
        self.assertEqual(
            IOCReviewStatus.objects.get(category='ip', value='198.51.100.7').note,
            'เครือข่ายภายนอก — ส่งคำขอ CVE-2020-25213',
        )
        self.assertTrue(TicketFieldChange.objects.filter(
            subtask=self.subtask, field_name='rca', new_value__contains='IOC Database',
        ).exists())

    def test_pushing_again_is_idempotent(self):
        self._add('ip', '198.51.100.7')
        rca_service.push_iocs(self.report, self.forensic)

        outcomes = rca_service.push_iocs(self.report, self.forensic)

        self.assertEqual(outcomes['created'], 0)
        self.assertEqual(outcomes['already_pushed'], 1)
        self.assertEqual(AnalystIOC.objects.filter(ioc_detail='198.51.100.7').count(), 1)

    def test_push_keeps_an_existing_annotation(self):
        IOCReviewStatus.objects.create(category='ip', value='198.51.100.7', note='FA note')
        self._add('ip', '198.51.100.7', note='from the report')

        rca_service.push_iocs(self.report, self.forensic)

        self.assertEqual(
            IOCReviewStatus.objects.get(category='ip', value='198.51.100.7').note, 'FA note',
        )

    def test_clean_normalizes_and_validates_values(self):
        indicator = RCAIndicator(rca=self.report, category='domain', value='  Evil.EXAMPLE. ')
        indicator.full_clean()
        self.assertEqual(indicator.value, 'evil.example')

        bad_hash = RCAIndicator(rca=self.report, category='hash', value='not-a-hash')
        with self.assertRaises(ValidationError):
            bad_hash.full_clean()

        bad_email = RCAIndicator(rca=self.report, category='email', value='nope')
        with self.assertRaises(ValidationError):
            bad_email.full_clean()


class RCAEditBookkeepingTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.forensic = _user('rca-edit-forensic', UserProfile.ROLE_FORENSIC)

    def test_an_edit_after_the_draft_marks_it_stale(self):
        subtask = _rca_request(_ticket(), self.forensic)
        report, _ = rca_service.get_or_create_rca(subtask, self.forensic)
        RCAReport.objects.filter(pk=report.pk).update(draft_generated_at=timezone.now())
        report.refresh_from_db()
        self.assertFalse(report.has_stale_draft)

        rca_service.record_edit(report, self.forensic, rca_service.SECTION_ROOT_CAUSES, 'RC-1')

        report.refresh_from_db()
        self.assertTrue(report.has_stale_draft)
        self.assertEqual(report.updated_by, self.forensic)

    def test_first_edit_starts_an_open_request_only_once(self):
        subtask = _rca_request(_ticket(), self.forensic)
        report, _ = rca_service.get_or_create_rca(subtask, self.forensic)

        rca_service.record_edit(report, self.forensic, rca_service.SECTION_ASSETS, 'a')
        rca_service.record_edit(report, self.forensic, rca_service.SECTION_ASSETS, 'b')

        subtask.refresh_from_db()
        self.assertEqual(subtask.status, TicketSubtask.STATUS_IN_PROGRESS)
        self.assertEqual(
            TicketFieldChange.objects.filter(subtask=subtask, field_name='status').count(), 1,
        )

    def test_an_edit_never_reopens_a_request_that_moved_on(self):
        subtask = _rca_request(_ticket(), self.forensic, status=TicketSubtask.STATUS_DONE)
        report, _ = rca_service.get_or_create_rca(subtask, self.forensic)

        rca_service.record_edit(report, self.forensic, rca_service.SECTION_ASSETS, 'a')

        subtask.refresh_from_db()
        self.assertEqual(subtask.status, TicketSubtask.STATUS_DONE)


RCA_TEMPLATE_PATH = (
    Path(settings.BASE_DIR) / 'apps' / 'incidents' / 'report_templates'
    / 'rca_report_template_v1.docx'
)


def _docx_placeholders(path):
    return {
        key
        for paragraph in _iter_paragraphs(Document(str(path)))
        for key in re.findall(r'\{\{([^}]+)\}\}', paragraph.text)
    }


class RCATemplateTest(TestCase):
    def test_build_script_matches_committed_template(self):
        script_path = Path(settings.BASE_DIR) / 'scripts' / 'build_rca_report_template_v1.py'
        spec = importlib.util.spec_from_file_location('build_rca_report_template_v1', script_path)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)

        with tempfile.TemporaryDirectory() as tmp:
            rebuilt_path = Path(tmp) / 'rebuilt_rca_template.docx'
            module.build(rebuilt_path)
            rebuilt = _docx_placeholders(rebuilt_path)

        self.assertEqual(rebuilt, _docx_placeholders(RCA_TEMPLATE_PATH))

    def test_template_carries_exactly_the_content_spec(self):
        expected = set()
        for kind, _label, spec in RCA_SECTION1_ROWS:
            if kind == 'kv':
                expected.add(spec)
            else:
                expected.update(key for key, _option in spec)
        for columns in RCA_REPEAT_TABLES.values():
            expected.update(key for _title, key, _width in columns if key)

        self.assertEqual(_docx_placeholders(RCA_TEMPLATE_PATH), expected)

    def test_forensic_types_match_the_model(self):
        self.assertEqual(
            [key for key, _label in RCA_FORENSIC_TYPES],
            [key for key, _label in RCAReport.FORENSIC_TYPE_CHOICES],
        )


def _docx_text(content):
    return '\n'.join(p.text for p in _iter_paragraphs(Document(BytesIO(content))))


class RCADraftTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.forensic = _user(
            'rca-draft-forensic', UserProfile.ROLE_FORENSIC,
            department='ปปกก.', phone='02-574-8209', first_name='วศิน', last_name='เชาว์บวร',
        )

    def _populated_rca(self):
        ticket = _ticket(
            incident_name='เว็บไซต์ถูกเปลี่ยนหน้า', severity='High',
            ncsa_severity=Ticket.NCSA_SEVERITY_NON_SEVERE, detailed_issue='Root Intrusion',
            asset_type='Server', is_emergency=True,
        )
        subtask = _rca_request(ticket, self.forensic)
        rca, _ = rca_service.get_or_create_rca(subtask, self.forensic)
        rca.forensic_types = ['host_disk', 'log_timeline', 'malware']
        rca.first_occurrence = _bkk(2023, 4, 21, 16, 0)
        rca.first_occurrence_note = 'โดยประมาณ'
        rca.scope_start = _bkk(2020, 11, 11, 0, 0)
        rca.scope_end = _bkk(2026, 9, 10, 0, 0)
        rca.save()
        RCAAsset.objects.create(
            rca=rca, order=1, host='ginfo', ip='111.111.11.111', detail='Joomla 1.5',
        )
        RCATimelineEntry.objects.create(
            rca=rca, occurred_at=_bkk(2023, 4, 21, 16, 27),
            occurred_until=_bkk(2023, 4, 21, 16, 38), host='NTDA',
            event='Brute Force MariaDB', evidence_file='authfail.log', evidence_line='14',
            excerpt='Access denied',
        )
        rc1 = RCARootCause.objects.create(
            rca=rca, order=1, category='Application', cause='Joomla 1.5 หมดการสนับสนุน',
        )
        rc2 = RCARootCause.objects.create(
            rca=rca, order=2, category='การตั้งค่าระบบ', cause='ไม่แยกบัญชีผู้ใช้',
        )
        rec = RCARecommendation.objects.create(rca=rca, order=1, action='ย้ายไปแพลตฟอร์มที่รองรับ')
        rec.root_causes.set([rc1, rc2])
        RCAIndicator.objects.create(
            rca=rca, category=RCAIndicator.CAT_HASH, value=HASH_A,
            label='wp-indos.php', source=RCAIndicator.SOURCE_ANALYST,
        )
        RCAIndicator.objects.create(
            rca=rca, category=RCAIndicator.CAT_IP, value='185.242.3.85',
            label='เครือข่ายภายนอก', note='ส่งคำขอ CVE', source=RCAIndicator.SOURCE_ANALYST,
        )
        RCAIndicator.objects.create(
            rca=rca, category=RCAIndicator.CAT_EMAIL, value='attacker@example.com',
            source=RCAIndicator.SOURCE_ANALYST,
        )
        return ticket, subtask, rca

    def test_draft_fills_data_and_leaves_no_placeholders(self):
        ticket, _subtask, rca = self._populated_rca()
        period = ticket.ticket_id[len('SOC-'):]

        filename, content = rca_service.generate_rca_draft(rca, self.forensic)

        self.assertEqual(filename, f'report_SOC-RCA-{period}_rca-v1.docx')
        text = _docx_text(content)
        self.assertNotIn('{{', text)
        self.assertIn(f'SOC-RCA-{period}', text)
        self.assertIn('เว็บไซต์ถูกเปลี่ยนหน้า', text)
        # Section 1 datetimes: first-occurrence carries its confirmable note, and
        # the scope renders as a Thai "start – end" range.
        self.assertIn('21 เม.ย. 2566 16:00 น. (โดยประมาณ)', text)
        self.assertIn('11 พ.ย. 2563 00:00 น. – 10 ก.ย. 2569 00:00 น.', text)
        self.assertIn('☑ Host / Disk Triage', text)
        self.assertIn('☐ Memory Forensic', text)
        self.assertIn('☑ High', text)          # SIEM severity
        self.assertIn('☑ สำคัญมาก', text)       # emergency → critical importance
        self.assertIn('☑ Server', text)
        # Evidence tables cloned per item.
        self.assertIn('ginfo', text)
        self.assertIn('Brute Force MariaDB', text)
        self.assertIn('16:27–16:38', text)
        self.assertIn(HASH_A, text)
        self.assertIn('185.242.3.85', text)
        self.assertIn('attacker@example.com', text)
        # Root-cause codes derived from order; recommendation cites them.
        self.assertIn('RC-1', text)
        self.assertIn('RC-2', text)
        self.assertIn('ย้ายไปแพลตฟอร์มที่รองรับ (RC-1, RC-2)', text)

    def test_draft_records_provenance_without_touching_ticket(self):
        ticket, _subtask, rca = self._populated_rca()

        _filename, content = rca_service.generate_rca_draft(rca, self.forensic)

        rca.refresh_from_db()
        self.assertEqual(rca.draft_template_version, 'rca-v1')
        self.assertEqual(rca.draft_generated_by, self.forensic)
        self.assertIsNotNone(rca.draft_generated_at)
        self.assertEqual(
            rca.draft_sha256, __import__('hashlib').sha256(content).hexdigest(),
        )
        ticket.refresh_from_db()
        self.assertEqual(ticket.report_template_version, '')
        self.assertEqual(ticket.report_sha256, '')
        self.assertIsNone(ticket.report_generated_by)

    def test_stale_flag_after_a_draft_then_edit(self):
        _ticket_obj, _subtask, rca = self._populated_rca()
        rca_service.generate_rca_draft(rca, self.forensic)
        rca.refresh_from_db()
        self.assertFalse(rca.has_stale_draft)

        rca_service.record_edit(rca, self.forensic, rca_service.SECTION_ASSETS, 'edited')

        rca.refresh_from_db()
        self.assertTrue(rca.has_stale_draft)

    def test_draft_of_an_empty_report_still_renders(self):
        subtask = _rca_request(_ticket(), self.forensic)
        rca, _ = rca_service.get_or_create_rca(subtask, self.forensic)
        rca.indicators.all().delete()
        rca.assets.all().delete()

        _filename, content = rca_service.generate_rca_draft(rca, self.forensic)

        self.assertNotIn('{{', _docx_text(content))
