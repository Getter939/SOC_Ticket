"""Analyst-uploaded IOC imports and the unified IOC Database view. No external lookup.

The IOC Database is the resting point for IOC data from two sources:
  • ticket IOCs  — entered by T1/T2 analysts on tickets (TicketIOC)
  • analyst IOCs — evidence the Forensic Analyst found externally and uploaded
                   here (AnalystIOC); NOT the MISP/TI-platform set

The FA reviews each indicator against MISP and sets a manual two-state status
(Checked / Not Checked), stored per (category, value) in IOCReviewStatus.

Import format (fixed five columns): ID, Category, File name, IOC detail, Note.
``ID`` is analyst-assigned and is the dedup identity — a re-imported ID is skipped
(the original kept). Each row is one indicator whose kind is ``Category`` and whose
value is ``IOC detail``; ``File name`` is context only.
"""

import csv
import io
import zipfile
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Count, Max
from django.utils import timezone
from openpyxl import load_workbook

from .ioc_values import TICKET_CATEGORY_CHOICES, normalize_for_category
from .models import AnalystIOC, AnalystIOCImport, IOCReviewStatus, TicketIOC

TICKET_CATEGORY_LABELS = dict(TICKET_CATEGORY_CHOICES)
SOURCE_LABELS = {'ticket': 'Ticket', 'analyst': 'FA'}

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 25 * 1024 * 1024
MAX_IMPORT_ROWS = 10000
MAX_REPORTED_ERRORS = 20

HEADERS = ('ID', 'Category', 'File name', 'IOC detail', 'Note')
COLUMNS = ('ext_id', 'category', 'file_name', 'ioc_detail', 'note')

# Header name (normalized) → column key. EN + TH.
ALIASES = {
    'id': 'ext_id', 'รหัส': 'ext_id',
    'category': 'category', 'ประเภท': 'category',
    'filename': 'file_name', 'ชื่อไฟล์': 'file_name',
    'iocdetail': 'ioc_detail', 'ioc': 'ioc_detail', 'detail': 'ioc_detail',
    'รายละเอียดioc': 'ioc_detail', 'รายละเอียด': 'ioc_detail',
    'note': 'note', 'หมายเหตุ': 'note',
}

# Category cell (normalized) → category code.
CATEGORY_ALIASES = {
    'hash': 'hash', 'sha256': 'hash', 'แฮช': 'hash',
    'ip': 'ip', 'ipaddress': 'ip', 'ที่อยู่ip': 'ip',
    'domain': 'domain', 'โดเมน': 'domain',
    'url': 'url', 'ยูอาร์แอล': 'url',
    'filepath': 'file_path', 'path': 'file_path', 'พาธไฟล์': 'file_path', 'พาธ': 'file_path',
}


def can_manage_inventory(user):
    profile = getattr(user, 'profile', None)
    return user.is_superuser or bool(profile and profile.is_forensic)


def validate_import_file(upload):
    if Path(upload.name).suffix.lower() not in {'.csv', '.xlsx'}:
        raise ValidationError('Upload a CSV or XLSX file.')
    if upload.size > MAX_IMPORT_BYTES:
        raise ValidationError('The import file must be 5 MB or smaller.')


def _key(value):
    return str(value or '').strip().lower().replace(' ', '').replace('-', '').replace('_', '')


def _read_rows(upload):
    validate_import_file(upload)
    upload.seek(0)
    content = upload.read(MAX_IMPORT_BYTES + 1)
    upload.seek(0)
    if len(content) > MAX_IMPORT_BYTES:
        raise ValidationError('The import file must be 5 MB or smaller.')
    if Path(upload.name).suffix.lower() == '.csv':
        try:
            yield from csv.reader(io.StringIO(content.decode('utf-8-sig')), strict=True)
        except (UnicodeError, csv.Error) as exc:
            raise ValidationError('Use a valid UTF-8 CSV file.') from exc
        return
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise ValidationError('The expanded workbook exceeds 25 MB.')
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            sheet = workbook.worksheets[0]
            if sheet.max_column and sheet.max_column > 5:
                raise ValidationError('Use only the five template columns.')
            sheet.reset_dimensions()
            for cells in sheet.iter_rows(max_col=6):
                if any(cell.data_type in {'f', 'e'} for cell in cells):
                    raise ValidationError('Formula and error cells are not allowed. Paste values first.')
                values = [cell.value for cell in cells]
                if values[-1] not in (None, ''):
                    raise ValidationError('Use only the five template columns.')
                yield values[:5]
        finally:
            workbook.close()
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError('Could not read this XLSX workbook. Use the five-column template.') from exc


def parse_import(upload):
    """Validate the whole file before persisting anything; report source row numbers.

    Row-value errors are collected and reported together (capped) so a file can be
    fixed in one pass; structural problems (bad header, wrong column count) fail fast.
    """
    rows = _read_rows(upload)
    try:
        try:
            header = next(rows)
        except StopIteration:
            raise ValidationError('The file is empty.')
        mapping = [ALIASES.get(_key(value)) for value in header]
        if len(mapping) != 5 or set(mapping) != set(COLUMNS):
            raise ValidationError('Required columns: ' + ', '.join(HEADERS))
        records = []
        errors = []
        for row_number, row in enumerate(rows, start=2):
            if row_number > MAX_IMPORT_ROWS + 1:
                raise ValidationError(f'Import at most {MAX_IMPORT_ROWS} rows at a time.')
            if not any(str(value).strip() for value in row if value is not None):
                continue
            if len(row) != 5:
                raise ValidationError(f'Row {row_number}: expected five columns.')
            values = dict(zip(mapping, (str(v).strip() if v is not None else '' for v in row)))
            try:
                values['category'] = CATEGORY_ALIASES.get(_key(values['category']))
                if not values['category']:
                    raise ValidationError('Category must be one of Hash, IP, Domain, URL, File Path.')
                if not values['ext_id']:
                    raise ValidationError('ID is required.')
                if len(values['ext_id']) > 100:
                    raise ValidationError('ID is too long (100 characters maximum).')
                if not values['ioc_detail']:
                    raise ValidationError('IOC detail is required.')
                values['ioc_detail'] = normalize_for_category(values['category'], values['ioc_detail'])
                if len(values['file_name']) > 255:
                    raise ValidationError('File name is too long (255 characters maximum).')
            except ValidationError as exc:
                errors.append(f'Row {row_number}: ' + '; '.join(exc.messages))
                if len(errors) >= MAX_REPORTED_ERRORS:
                    errors.append(f'More rows have errors; showing the first {MAX_REPORTED_ERRORS}.')
                    break
                continue
            records.append(values)
        if errors:
            raise ValidationError(errors)
        if not records:
            raise ValidationError('The file contains no IOC records.')
        return records
    finally:
        rows.close()


def import_inventory(upload, user):
    records = parse_import(upload)
    batch = AnalystIOCImport(original_name=Path(upload.name).name, uploaded_by=user)
    try:
        with transaction.atomic():
            batch.file.save(batch.original_name, upload, save=False)
            batch.row_count = len(records)
            batch.save()
            for values in records:
                record, added = AnalystIOC.objects.get_or_create(
                    ext_id=values['ext_id'],
                    defaults={**values, 'source_import': batch, 'last_seen_import': batch},
                )
                if added:
                    batch.added_count += 1
                else:
                    # Same ID = duplicate: keep the original record. Only advance
                    # last_seen_import when this batch is at least as new, so an
                    # older/slower import can't regress it.
                    AnalystIOC.objects.filter(
                        pk=record.pk, last_seen_import__uploaded_at__lte=batch.uploaded_at,
                    ).update(last_seen_import=batch)
                    batch.skipped_count += 1
            batch.save(update_fields=['added_count', 'skipped_count'])
    except Exception:
        # File storage does not participate in a database rollback.
        if batch.file.name:
            batch.file.delete(save=False)
        raise
    return batch


def build_ioc_database(query='', status='all', source='all', category='all'):
    """Every distinct indicator (category, value) from BOTH sources, annotated with
    the FA's manual review status. Returns (rows, counts).

      • sources      — 'ticket' (TicketIOC) and/or 'analyst' (active AnalystIOC)
      • status       — 'checked' / 'not_checked' from IOCReviewStatus (default not_checked)

    ``counts`` are over the whole database (before filters), for the toolbar summary.
    """
    observed = {}
    for row in (TicketIOC.objects.values('category', 'value')
                .annotate(ticket_count=Count('ticket', distinct=True),
                          last_seen=Max('ticket__created_at'))):
        observed[(row['category'], row['value'])] = row

    analyst = {}
    for rec in AnalystIOC.objects.filter(is_active=True).select_related('last_seen_import'):
        analyst[(rec.category, rec.ioc_detail)] = rec

    checked_keys = {
        (s['category'], s['value'])
        for s in IOCReviewStatus.objects.filter(checked=True).values('category', 'value')
    }

    rows = []
    for key in set(observed) | set(analyst):
        category_code, value = key
        obs = observed.get(key)
        rec = analyst.get(key)
        sources = ([source_ for source_, present in
                    (('ticket', obs is not None), ('analyst', rec is not None)) if present])
        if obs:
            last_seen = obs['last_seen']
        elif rec is not None:
            last_seen = rec.last_seen_import.uploaded_at if rec.last_seen_import_id else rec.created_at
        else:
            last_seen = None
        rows.append({
            'category': category_code,
            'category_label': TICKET_CATEGORY_LABELS.get(category_code, category_code),
            'value': value,
            'sources': sources,
            'source_label': ' + '.join(SOURCE_LABELS[s] for s in sources),
            'status': 'checked' if key in checked_keys else 'not_checked',
            'ticket_count': obs['ticket_count'] if obs else 0,
            'last_seen': last_seen,
            'analyst_rec': rec,
        })

    counts = {
        'total': len(rows),
        'checked': sum(1 for row in rows if row['status'] == 'checked'),
        'not_checked': sum(1 for row in rows if row['status'] == 'not_checked'),
    }

    query = (query or '').strip().lower()
    if query:
        def matches(row):
            rec = row['analyst_rec']
            return (query in row['value'].lower()
                    or (rec is not None and (query in (rec.ext_id or '').lower()
                                             or query in (rec.note or '').lower())))
        rows = [row for row in rows if matches(row)]
    if status in ('checked', 'not_checked'):
        rows = [row for row in rows if row['status'] == status]
    if source in ('ticket', 'analyst'):
        rows = [row for row in rows if source in row['sources']]
    if category in TICKET_CATEGORY_LABELS:
        rows = [row for row in rows if row['category'] == category]

    # Not-checked first (the worklist), then busiest, then value.
    rows.sort(key=lambda row: (row['status'] != 'not_checked', -row['ticket_count'], row['value']))
    return rows, counts


def set_review_status(category, value, checked, user):
    """Set the FA's manual Checked / Not-Checked status for one indicator."""
    value = (value or '').strip()[:500]
    if not value or category not in TICKET_CATEGORY_LABELS:
        raise ValidationError('Unknown IOC.')
    status, _ = IOCReviewStatus.objects.update_or_create(
        category=category, value=value,
        defaults={'checked': bool(checked), 'updated_by': user},
    )
    return status


def remove_analyst_ioc(pk, user):
    """Soft-remove a mistakenly uploaded analyst IOC (hidden from the database)."""
    return AnalystIOC.objects.filter(pk=pk, is_active=True).update(
        is_active=False, removed_by=user, removed_at=timezone.now())
