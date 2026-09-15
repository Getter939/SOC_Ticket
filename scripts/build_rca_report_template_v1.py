"""Build the Root Cause Analysis (RCA) report DOCX template.

Mirrors the Forensic Analyst's own report (``Report_SOC-RCA-…_v5.odt``): A4, a
charcoal title block with a grey case-number band, dark section bands, grey
table headers, the case number and TLP:AMBER marker in every page header, and a
confidentiality footer with page numbers.

The document is a *template*, filled per request by ``apps.incidents.rca_report``:
  * Section 1 and the page header through ``{{placeholder}}`` substitution;
  * every repeatable table (sections 4, 5, 6.2, 7.x, 8) has a header row plus ONE
    prototype row of ``{{prefix_field}}`` placeholders, cloned once per item.
Narrative sections carry bracketed guidance text for the analyst to replace in
Word — no placeholders, so nothing is filled there.

Layout data (titles, rows, columns, guidance) lives in apps/incidents/rca_content.py.

Run from the repo root:  python scripts/build_rca_report_template_v1.py
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

from scripts.build_report_template_v2 import (  # noqa: E402
    BODY_FONT,
    SYMBOL_FONT,
    add_runs,
    clear_table_borders,
    set_cell_margins,
    set_cell_shading,
    set_run_font,
    set_table_borders,
)
from apps.incidents.rca_content import (  # noqa: E402
    RCA_ACCOUNT_COLUMNS,
    RCA_ASSET_COLUMNS,
    RCA_CASE_LABEL,
    RCA_FIGURES,
    RCA_FILE_BEHAVIOR_COLUMNS,
    RCA_FILE_PATH_COLUMNS,
    RCA_FOOTER_LEFT,
    RCA_GUIDANCE,
    RCA_HASH_COLUMNS,
    RCA_HEADER_LEFT,
    RCA_IP_COLUMNS,
    RCA_PRIMARY_ROOT_CAUSE_LABEL,
    RCA_RECOMMENDATION_COLUMNS,
    RCA_ROOT_CAUSE_COLUMNS,
    RCA_SECTION1_ROWS,
    RCA_SECTION_TITLES,
    RCA_SMALL_TEXT_KEYS,
    RCA_SUBSECTION_TITLES,
    RCA_SUBTITLE,
    RCA_TIMELINE_COLUMNS,
    RCA_TITLE,
    RCA_TLP_LABEL,
    RCA_WEB_COLUMNS,
)

TEMPLATES_DIR = BASE_DIR / 'apps' / 'incidents' / 'report_templates'
OUTPUT_PATH = TEMPLATES_DIR / 'rca_report_template_v1.docx'
TLP_ICON_PATH = TEMPLATES_DIR / 'assets' / 'tlp_amber.png'

# Palette sampled from the analyst's ODT.
BAND_DARK = '404040'
HEAD_GRAY = '737373'
CASE_BAND = 'D9D9D9'
LABEL_BG = 'F7F7F7'
BORDER = 'BFBFBF'
TEXT = '262626'
MUTED = '595959'
GUIDE = '7F7F7F'
WHITE = 'FFFFFF'
TLP_AMBER = 'FFC000'

CONTENT_CM = 17.0
BODY_PT = 14
TABLE_PT = 12
SMALL_PT = 10.5


# ── helpers ────────────────────────────────────────────────────────────── #

def _first(cell):
    paragraph = cell.paragraphs[0]
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    return paragraph


def _widths(table, widths_cm, *, top_align=False):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, width in enumerate(widths_cm):
        table.columns[index].width = Cm(width)
    align = WD_CELL_VERTICAL_ALIGNMENT.TOP if top_align else WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in table.rows:
        for index, width in enumerate(widths_cm):
            cell = row.cells[index]
            cell.width = Cm(width)
            cell.vertical_alignment = align
            set_cell_margins(cell, top=30, start=80, bottom=30, end=80)


def _spacer(doc, points=6):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(points)


def _shade_run(run, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    run._element.get_or_add_rPr().append(shd)


def _add_field(paragraph, instruction, size, color):
    """Append a Word field (PAGE / NUMPAGES) as begin/instr/separate/result/end runs."""
    def field_run(child):
        run = paragraph.add_run()
        set_run_font(run, size=size, color=color)
        run._r.append(child)

    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    field_run(begin)
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = f' {instruction} '
    field_run(instr)
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    field_run(separate)
    result = paragraph.add_run('1')
    set_run_font(result, size=size, color=color)
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    field_run(end)


def _repeat_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement('w:tblHeader'))


def _checks(options, size):
    segments = []
    for index, (key, label) in enumerate(options):
        if index:
            segments.append(('   ', BODY_FONT, size, TEXT, False))
        segments.append((f'{{{{{key}}}}}', SYMBOL_FONT, size - 1, TEXT, False))
        segments.append((f' {label}', BODY_FONT, size, TEXT, False))
    return segments


# ── page furniture ─────────────────────────────────────────────────────── #

def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = doc.styles['Normal']
    normal.font.name = BODY_FONT
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        normal._element.rPr.rFonts.set(qn(attr), BODY_FONT)
    normal.font.size = Pt(BODY_PT)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0


def build_header(doc):
    header = doc.sections[0].header
    header.is_linked_to_previous = False
    table = header.add_table(rows=1, cols=2, width=Cm(CONTENT_CM))
    clear_table_borders(table)
    _widths(table, [10.5, 6.5])
    left, right = table.rows[0].cells
    add_runs(_first(left), [(RCA_HEADER_LEFT, BODY_FONT, 10, MUTED, False)])
    paragraph = _first(right)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_runs(paragraph, [('{{rca_case_no}}', BODY_FONT, 10, MUTED, False), ('  ', BODY_FONT, 10, MUTED, False)])
    badge = paragraph.add_run(f' {RCA_TLP_LABEL} ')
    set_run_font(badge, size=10, color=TLP_AMBER, bold=True)
    _shade_run(badge, '000000')
    if TLP_ICON_PATH.exists():
        paragraph.add_run(' ').add_picture(str(TLP_ICON_PATH), height=Cm(0.42))
    empty = header.paragraphs[0]
    empty._p.getparent().remove(empty._p)


def build_footer(doc):
    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    table = footer.add_table(rows=1, cols=2, width=Cm(CONTENT_CM))
    clear_table_borders(table)
    _widths(table, [12.5, 4.5])
    left, right = table.rows[0].cells
    add_runs(_first(left), [(RCA_FOOTER_LEFT, BODY_FONT, 10, MUTED, False)])
    paragraph = _first(right)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_runs(paragraph, [('หน้า ', BODY_FONT, 10, MUTED, False)])
    _add_field(paragraph, 'PAGE', 10, MUTED)
    add_runs(paragraph, [(' / ', BODY_FONT, 10, MUTED, False)])
    _add_field(paragraph, 'NUMPAGES', 10, MUTED)
    empty = footer.paragraphs[0]
    empty._p.getparent().remove(empty._p)


# ── content blocks ─────────────────────────────────────────────────────── #

def add_title(doc):
    table = doc.add_table(rows=2, cols=1)
    clear_table_borders(table)
    _widths(table, [CONTENT_CM])
    title_cell, case_cell = table.rows[0].cells[0], table.rows[1].cells[0]
    set_cell_shading(title_cell, BAND_DARK)
    set_cell_margins(title_cell, top=110, bottom=90)
    title = _first(title_cell)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(title, [(RCA_TITLE, BODY_FONT, 20, WHITE, True)])
    subtitle = title_cell.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(subtitle, [(RCA_SUBTITLE, BODY_FONT, 13, WHITE, True)])
    set_cell_shading(case_cell, CASE_BAND)
    set_cell_margins(case_cell, top=50, bottom=50)
    case = _first(case_cell)
    case.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(case, [
        (f'{RCA_CASE_LABEL}  ', BODY_FONT, 12, MUTED, False),
        ('{{rca_case_no}}', BODY_FONT, 13, TEXT, True),
    ])
    _spacer(doc, 10)


def add_band(doc, number):
    thai, english = RCA_SECTION_TITLES[number]
    table = doc.add_table(rows=1, cols=1)
    clear_table_borders(table)
    _widths(table, [CONTENT_CM])
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, BAND_DARK)
    set_cell_margins(cell, top=50, bottom=50, start=100, end=100)
    paragraph = _first(cell)
    paragraph.paragraph_format.keep_with_next = True
    add_runs(paragraph, [
        (f'{number}.  {thai}', BODY_FONT, 16, WHITE, True),
        (f'   ({english})', BODY_FONT, 11, CASE_BAND, False),
    ])
    _spacer(doc, 4)


def add_subheading(doc, number):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    add_runs(paragraph, [(f'{number} {RCA_SUBSECTION_TITLES[number]}', BODY_FONT, BODY_PT, TEXT, True)])


def add_guidance(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.first_line_indent = Cm(1.0)
    run = paragraph.add_run(text)
    set_run_font(run, size=BODY_PT, color=GUIDE)
    run.italic = True


def add_label(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.first_line_indent = Cm(1.0)
    paragraph.paragraph_format.keep_with_next = True
    add_runs(paragraph, [(text, BODY_FONT, BODY_PT, TEXT, True)])


def add_section1(doc):
    table = doc.add_table(rows=len(RCA_SECTION1_ROWS), cols=2)
    _widths(table, [5.2, 11.8])
    set_table_borders(table, color=BORDER, size='4')
    for index, (kind, label, spec) in enumerate(RCA_SECTION1_ROWS):
        left, right = table.rows[index].cells
        set_cell_shading(left, LABEL_BG)
        add_runs(_first(left), [(label, BODY_FONT, 13, TEXT, False)])
        if kind == 'kv':
            add_runs(_first(right), [(f'{{{{{spec}}}}}', BODY_FONT, 13, TEXT, False)])
        else:
            add_runs(_first(right), _checks(spec, 13))
    _spacer(doc, 10)


def add_repeat_table(doc, columns):
    """Header row (repeated on every page) + one prototype row of placeholders.

    A column whose key is None stays empty for the analyst to fill in Word.
    """
    table = doc.add_table(rows=2, cols=len(columns))
    _widths(table, [width for _title, _key, width in columns], top_align=True)
    set_table_borders(table, color=BORDER, size='4')
    header, prototype = table.rows
    _repeat_header_row(header)
    for index, (title, key, _width) in enumerate(columns):
        head_cell = header.cells[index]
        set_cell_shading(head_cell, HEAD_GRAY)
        head_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        add_runs(_first(head_cell), [(title, BODY_FONT, TABLE_PT, WHITE, True)])
        if key:
            size = SMALL_PT if key in RCA_SMALL_TEXT_KEYS else TABLE_PT
            add_runs(_first(prototype.cells[index]), [(f'{{{{{key}}}}}', BODY_FONT, size, TEXT, False)])
    _spacer(doc, 10)


def add_figure_slot(doc, placeholder, caption):
    table = doc.add_table(rows=1, cols=1)
    _widths(table, [CONTENT_CM])
    set_table_borders(table, color=BORDER, size='4')
    cell = table.rows[0].cells[0]
    set_cell_margins(cell, top=600, bottom=600)
    paragraph = _first(cell)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(placeholder)
    set_run_font(run, size=BODY_PT, color=GUIDE)
    run.italic = True
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(8)
    add_runs(caption_paragraph, [(caption, BODY_FONT, 12, TEXT, True)])


# ── document ───────────────────────────────────────────────────────────── #

def build(output_path=OUTPUT_PATH):
    doc = Document()
    style_doc(doc)
    build_header(doc)
    build_footer(doc)
    add_title(doc)

    add_band(doc, '1')
    add_section1(doc)

    add_band(doc, '2')
    add_guidance(doc, RCA_GUIDANCE['2'])
    add_label(doc, RCA_PRIMARY_ROOT_CAUSE_LABEL)
    add_guidance(doc, RCA_GUIDANCE['2_primary'])
    _spacer(doc, 8)

    add_band(doc, '3')
    for placeholder, caption in RCA_FIGURES:
        add_figure_slot(doc, placeholder, caption)

    add_band(doc, '4')
    add_repeat_table(doc, RCA_ASSET_COLUMNS)

    add_band(doc, '5')
    add_guidance(doc, RCA_GUIDANCE['5'])
    add_repeat_table(doc, RCA_TIMELINE_COLUMNS)

    add_band(doc, '6')
    add_subheading(doc, '6.1')
    add_guidance(doc, RCA_GUIDANCE['6.1'])
    add_subheading(doc, '6.2')
    add_repeat_table(doc, RCA_ROOT_CAUSE_COLUMNS)

    add_band(doc, '7')
    add_subheading(doc, '7.1')
    add_repeat_table(doc, RCA_FILE_PATH_COLUMNS)
    add_subheading(doc, '7.2')
    add_repeat_table(doc, RCA_FILE_BEHAVIOR_COLUMNS)
    add_guidance(doc, RCA_GUIDANCE['7.2'])
    add_subheading(doc, '7.3')
    add_repeat_table(doc, RCA_IP_COLUMNS)
    add_subheading(doc, '7.4')
    add_repeat_table(doc, RCA_WEB_COLUMNS)
    add_subheading(doc, '7.5')
    add_repeat_table(doc, RCA_HASH_COLUMNS)
    add_subheading(doc, '7.6')
    add_repeat_table(doc, RCA_ACCOUNT_COLUMNS)

    add_band(doc, '8')
    add_guidance(doc, RCA_GUIDANCE['8'])
    add_repeat_table(doc, RCA_RECOMMENDATION_COLUMNS)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(output_path)


if __name__ == '__main__':
    build()
